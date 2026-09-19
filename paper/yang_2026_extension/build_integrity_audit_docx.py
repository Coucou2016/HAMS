from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_manuscript_docx import (
    INK,
    MUTED,
    add_body_paragraph,
    add_markdown_table,
    clean_inline,
    configure_document,
    set_font,
)


PAPER_DIR = Path(__file__).resolve().parent
SOURCE = PAPER_DIR / "RESEARCH_INTEGRITY_AUDIT_CN.md"
OUTPUT = PAPER_DIR / "Research_Integrity_Audit_CN_TheoryFramework.docx"


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Pt(14)
    paragraph.paragraph_format.right_indent = Pt(8)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("\n".join(lines))
    set_font(run, "Consolas", 8.3)


def build_document(source: Path, output: Path, header_text: str, subject: str, *, landscape: bool = False) -> Path:
    doc = Document()
    configure_document(doc)
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        doc.styles["Normal"].font.size = Pt(9.6)
        doc.styles["Normal"].paragraph_format.line_spacing = 1.05
        doc.styles["Title"].font.size = Pt(17)
        doc.styles["Title"].paragraph_format.space_before = Pt(26)
        doc.styles["Title"].paragraph_format.space_after = Pt(18)
    header = section.header.paragraphs[0]
    header.text = header_text
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(header.runs[0], "SimSun", 8.0, color=INK)

    normal = doc.styles["Normal"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    for heading_name in ["Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[heading_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")

    lines = source.read_text(encoding="utf-8").splitlines()
    title = clean_inline(lines[0].lstrip("# "))
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(26 if landscape else 36)
    paragraph.paragraph_format.space_after = Pt(18 if landscape else 22)
    run = paragraph.add_run(title)
    set_font(run, "SimHei", 18, bold=True, color=INK)

    index = 1
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        text = " ".join(paragraph_buffer)
        paragraph_buffer.clear()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(clean_inline(text))
        set_font(run, "SimSun", 10.5)

    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            doc.add_paragraph(clean_inline(stripped[3:]), style="Heading 1")
            index += 1
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(clean_inline(stripped[4:]), style="Heading 2")
            index += 1
            continue
        bold_lead = re.match(r"^\*\*(.+?)\*\*\s*(.*)$", stripped)
        if bold_lead:
            flush_paragraph()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            label = paragraph.add_run(clean_inline(bold_lead.group(1)))
            set_font(label, "SimSun", 10.0, bold=True)
            value = paragraph.add_run(" " + clean_inline(bold_lead.group(2)))
            set_font(value, "SimSun", 10.0)
            index += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            add_code_block(doc, code_lines)
            index += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            if len(rows) > 1 and all(set(cell) <= {"-", ":"} for cell in rows[1]):
                rows.pop(1)
            add_markdown_table(doc, rows)
            continue
        if re.match(r"^[-*]\s+", stripped):
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Bullet")
            run = paragraph.add_run(clean_inline(stripped[2:]))
            set_font(run, "SimSun", 10.3)
            index += 1
            continue
        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    properties = doc.core_properties
    properties.title = title
    properties.subject = subject
    properties.author = "Sea-landing dynamics research project"
    properties.keywords = "data lineage, reproducibility, research integrity, reviewer response"
    doc.save(output)
    print(output)
    return output


def build() -> Path:
    return build_document(
        SOURCE,
        OUTPUT,
        "海上回收着陆论文 | 真实性、准确性与完整性审查",
        "Research integrity and reproducibility audit",
        landscape=True,
    )


if __name__ == "__main__":
    build()
