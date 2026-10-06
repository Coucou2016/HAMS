"""Refine the supplied author manuscript without replacing its math or figures."""
from __future__ import annotations

import difflib
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "20260925-审稿意见"
PACKAGE = INPUT / "HAMS_论文修订_审查证据_候选代码补丁_20260926" / "HAMS_review_20260926"
SOURCE_DOCX = INPUT / "HAMS_manuscript_revised.docx"
SOURCE_MD = PACKAGE / "manuscript" / "manuscript_revised_pending_contact_rerun.md"
OUTPUT = Path(__file__).resolve().parent / "revision_20260927"

# Whole-paragraph edits are restricted to prose without embedded mathematics.
PROSE = [
    (
        "Landing on a floating platform depends",
        "Landing on a floating platform depends on motion at the contact points, where platform rotation contributes to both displacement and velocity. A partitioned formulation combines frequency-domain wave-body interaction, finite-band radiation memory and independent four-foot contact. Complex response operators transfer platform motion to the landing footprint, and the contact wrench supplies the platform feedback. A reduced comparison with the drop tests of Yang et al. gives buffer-stroke peak errors below 4.1% for two asymmetric sequences using parameters fitted only to simultaneous touchdown. Acceleration and force histories remain discrepant, while 127 of 128 multi-start fits are near-optimal, indicating poor practical identifiability. For a prescribed 120 m by 50 m barge, the sampled heave response reaches 1.113 m/m. The selected medium-to-fine mesh change is 5.88%, above the 5% criterion. Screening with 1000 random-wave realizations per condition and bootstrap intervals places 41 of 63 conditions within a preliminary linear-motion envelope. These statistics describe finite-record deck-motion extremes, not landing success. The archived full-scale calculation meets a 2% fixed-point update criterion after four passes but fails contact-refinement checks. Its inertia mapping, collision-deck reference and executed load transfer also require correction. The corresponding landing histories are retained solely as pre-correction diagnostics. Independent subproblem reruns support computational reproducibility within the stated scope; corrected coupled integration and renewed convergence tests are required before interpreting the full-scale contact loads or touchdown sequence physically.",
        "摘要按方法、分层结果和适用范围组织，保留失败结论及全部关键数值。",
    ),
    (
        "This study examines that interface",
        "This study develops and assesses the interface between linear platform hydrodynamics and independent four-foot contact. The formulation transfers complex platform response operators to the landing center and footprint, resolves separate foot-contact states, and returns vertical force and roll/pitch moments to the same platform operator. Its contribution lies in the motion and load interfaces, case construction and numerical assessment; the boundary-element, radiation-memory and rigid-body formulations are established methods. The reduced 5.2 t drop-test comparison and the full-scale example are evaluated separately because they use different vehicle models and parameters. Powered descent, guidance, propellant sloshing, dynamic positioning and nonlinear free-surface effects are outside the present scope. The full-scale example is retained as a diagnostic pending the corrections described in Sections 2.4-2.5.",
        "用研究框架界定贡献，软件实现归属留在方法和数据可用性部分。",
    ),
    (
        "The archived main calculation transfers",
        "The archived main calculation uses pointwise linear interpolation (`np.interp`) to transfer the sampled wrench. The interval-averaging routine elsewhere in the repository is not on this execution path, so the archived results do not use the conservative transfer previously described. The candidate correction integrates the piecewise-linear fine-grid wrench over control intervals bounded by target-time midpoints and divides each integral by its interval width. With identical source and target endpoints, the half-width endpoint intervals ensure that the target-grid trapezoidal impulse equals the source-grid impulse for each wrench component. This preserves resultant impulse under the stated quadrature; peak-load accuracy, interface work and platform response still require separate checks. The candidate transfer has not yet produced a corrected coupled landing history.",
        "补足守恒重采样的同覆盖区间条件，区分算法性质与已执行路径。",
    ),
    (
        "Equal-and-opposite arrays are constructed",
        "The wrench reducer constructs equal-and-opposite arrays from the same foot-force samples. A zero residual therefore checks sign and reference-point assembly rather than independently verifying the deck reaction. An independent comparison requires reactions recorded on the second body or a separate benchmark. Contact power also requires force, moment and velocity to refer to the same point on each body. Pairing a platform-origin moment with the vehicle center-of-gravity velocity does not satisfy that requirement, particularly when the forces act on articulated feet. The mixed-reference work proxy is excluded from the energy evidence; Figure 1 depicts the intended interface rather than certifying the archived implementation.",
        "用同点功率关系解释审查发现，删除重复的答辩式强调。",
    ),
    (
        "The independent review rebuilt",
        "A separate reproducibility assessment rebuilt the supplied Fortran source, ran the configured Cylinder case, and evaluated the medium barge at 0.6, 0.65, 0.775, 0.8 and 1.0 rad/s for seven headings. At these frequencies, the selected diagonal hydrodynamic coefficients match the archived medium outputs at file precision, and the principal motion amplitudes differ only at roundoff level. Some very small coupling entries differ between runs. The Cylinder comparison retains a substantive damping discrepancy against the supplied benchmark, quantified in Section 4.2. These executions establish reproducibility for the specified outputs and do not complete solver certification.",
        "明确复算覆盖范围，将评价放到对应结果节。",
    ),
    (
        "The review also independently integrated",
        "The same assessment integrated the three reduced drop-test cases, repeated the simultaneous-case fit, reconstructed 1000 phase realizations for one critical sea condition, and replayed the platform equation under archived contact forcing. The supplied execution records identify the inputs, environment and comparisons for these checks. That separate review environment lacked PyChrono, and no corrected full-scale contact history is supplied. Platform replay under an existing wrench therefore provides a subproblem check, not a new coupled landing calculation. No experiment for the complete modeled barge-vehicle combination is available for validation.",
        "将外部审查环境与当前工作机区分，避免把所给日志写成本轮新运行。",
    ),
    (
        "Figure 6 and Table 4 retain",
        "Figure 6 and Table 4 document the archived pre-correction calculation. Wave, plume and contact forcing share the 120 m by 50 m hydrodynamic operator, but the inertia mapping, collision-deck reference and load-transfer implementation require the corrections identified in Section 2. The values below characterize that archived implementation and provide a baseline for subsequent comparison. They do not constitute corrected physical contact predictions.",
        "集中声明旧接触结果身份，保留数据和待重算状态。",
    ),
    (
        "The stored multibody response contains",
        "The stored multibody response contains 3638 samples, compared with 40,001 nominal contact-grid samples over 20 s. Its median spacing is 0.0055 s rather than the 0.0005 s integration step. Integrating the compact history gives a positive vertical impulse approximately 0.132% different from the archived full-grid summary. That comparison assesses one integral and does not bound the error in peaks or event times omitted during compression. The candidate recorder writes all recorded per-pass integration nodes separately from plotting histories. Until corrected dynamics and those records are available, the archived sequence remains a numerical diagnostic.",
        "区分积分节点、压缩记录和汇总值，保留不可恢复高频信息的限制。",
    ),
    (
        "The immediate revision can remain",
        "A corrected assessment requires consistent inertia axes and deck references, conservative resultant-load transfer, and complete recorded contact histories. Production, step-refinement and stiffness cases must satisfy the same interface-closure criterion, with impulse comparisons restricted to a common time window. Figure 6, Table 4 and the associated contact claims can then be regenerated from one corrected output set. The wave-screening and reduced-test results retain their separate evidence and limitations; they cannot substitute for this coupled calculation.",
        "把工作安排式语言改为数值评估要求，保留最小复算顺序。",
    ),
    (
        "A partitioned formulation links",
        "The partitioned formulation links frequency-domain wave-body response, finite-band radiation memory and independent foot-contact states through local deck kinematics and a reference-point wrench. Transforming complex response operators preserves phase relationships across the landing footprint, while wrench reduction includes the moment of offset contact forces. These interfaces define the proposed calculation; their predictive accuracy must be assessed separately from the reproducibility of each subproblem.",
        "结论先陈述方法功能，再说明证据对应的判断。",
    ),
    (
        "The archived landing satisfies its final",
        "The archived landing meets its final update criterion after four passes, but contact-refinement and stable-standing criteria remain unmet. Inertia-axis mapping, collision-deck reference and load transfer introduce additional implementation errors that require a complete corrected rerun. The supplied candidate changes have analytical tests but no corrected contact history. Consequently, the reported contact forces, strokes and touchdown sequence remain pre-correction diagnostics and do not qualify design loads or operational limits.",
        "收束结论，保留未修复实现错误及未完成重算的实质限定。",
    ),
    (
        "The supplied archive contains case inputs",
        "The project repository and accompanying review package contain case inputs, meshes, locally generated hydrodynamic outputs, deck-point response operators, realization-level wave statistics, digitized reference targets, compact multibody histories, refinement summaries and figure-generation code. The main archived multibody history is downsampled. The review package separately records the specified subproblem reruns, file hashes, comparison scripts and candidate source changes. File identity checks connect those records to their inputs; they do not establish experimental authenticity or physical accuracy.",
        "将数据可用性改成可核查的材料清单，避免自证式表述。",
    ),
    (
        "Published absorber curves, drop-test traces",
        "Published absorber curves and vehicle parameters are attributed model inputs; digitized drop-test traces are external comparison targets. Model-generated responses are stored separately. HAMS and Project Chrono are third-party solver dependencies, while the present implementation supplies the case configuration, interfaces and analyses described here. The absence of PyChrono in the separate review environment limits that assessment, not the availability of the source code. No corrected coupled history or renewed contact-convergence evidence is included in this manuscript version.",
        "保留第三方求解器和参考材料归属，消除环境范围歧义。",
    ),
]

FRAGMENTS = [
    ("A Partitioned Hydromechanical Framework for Four Leg Reusable Launch Vehicle Landing on a Floating Barge", "A Partitioned Hydromechanical Framework for Four-Leg Reusable Launch Vehicle Landing on a Floating Barge", "标题复合修饰语统一。"),
    ("are returned to the present platform equation because horizontal restoring data are unavailable.", "are returned to the present platform equation under the imposed surge, sway and yaw constraints.", "图1图注与正文一致：缺少恢复力数据不等于水平自由度物理静止。"),
    ("Failing entries remain failures even when neighboring curves appear smooth.", "The criterion is applied to the selected numerical entries rather than inferred from curve smoothness.", "保留判据，调整审稿答辩式语气。"),
    ("The latter is not rounded down to a pass.", "The medium-to-fine comparison therefore does not meet the stated criterion.", "直接陈述网格比较结论。"),
    ("### 4.3 Random-wave motion at the actual foot locations", "### 4.3 Random-wave motion at the modeled foot locations", "支腿布置是模型输入，不是真实型号测量几何。"),
    ("Actual feet use the 6.926 m-radius layout;", "Modeled feet use the 6.926 m-radius layout;", "图5与代理几何定义一致。"),
    ("The source values are unchanged; labels identify their actual meaning. No panel represents a corrected Chrono rerun.", "All panels retain pre-correction values; corrected coupled results remain pending.", "图6保留明确状态，压缩重复措辞。"),
    ("The ensemble demonstrates non-uniqueness of the fitted physical parameters.", "The ensemble shows non-uniqueness of the fitted effective model parameters.", "图S1不将观测映射比等有效系数称为独立物理部件参数。"),
]


def normalized(text: str) -> str:
    return text.replace("\u00a0", " ").replace("’", "'").replace("‘", "'")


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def visible_text(paragraph) -> str:
    return "".join(node.text or "" for node in paragraph._p.xpath(".//w:t"))


def replace_text(paragraph, old: str, new: str) -> None:
    """Edit only w:t nodes; OMML, drawings, formatting and relationships remain."""
    nodes = paragraph._p.xpath(".//w:t")
    plain = "".join(node.text or "" for node in nodes)
    norm = normalized(plain)
    if norm.count(normalized(old)) != 1:
        raise ValueError(f"Ambiguous text: {old[:80]}")
    start = norm.index(normalized(old))
    end = start + len(old)
    offset = 0
    inserted = False
    for node in nodes:
        value = node.text or ""
        stop = offset + len(value)
        if offset < end and stop > start:
            left = value[:max(0, start - offset)]
            right = value[max(0, end - offset):]
            node.text = left + (new if not inserted else "") + right
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            inserted = True
        offset = stop
    if normalized(visible_text(paragraph)) != norm[:start] + normalized(new) + norm[end:]:
        raise AssertionError("Text replacement failed")


def structural_record(path: Path) -> dict:
    doc = Document(path)
    with zipfile.ZipFile(path) as package:
        media = {name: hashlib.sha256(package.read(name)).hexdigest()
                 for name in package.namelist() if name.startswith("word/media/")}
    def digest(nodes):
        return [hashlib.sha256(etree.tostring(node, method="c14n")).hexdigest() for node in nodes]
    return {"tables": digest([table._tbl for table in doc.tables]),
            "math": digest(doc.element.xpath("//m:oMath")),
            "media": media, "inline_shapes": len(doc.inline_shapes)}


def refine() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    source_hash = sha(SOURCE_DOCX)
    doc = Document(SOURCE_DOCX)
    original_md = SOURCE_MD.read_text(encoding="utf-8")
    markdown = original_md
    changes = []
    for prefix, new, reason in PROSE:
        paragraphs = [p for p in doc.paragraphs if normalized(visible_text(p)).startswith(prefix)]
        blocks = [b for b in markdown.split("\n\n") if b.startswith(prefix)]
        if len(paragraphs) != 1 or len(blocks) != 1:
            raise ValueError(f"Source paragraph not unique: {prefix}")
        paragraph, old_md = paragraphs[0], blocks[0]
        if paragraph._p.xpath(".//m:oMath | .//w:drawing"):
            raise ValueError(f"Whole-paragraph edit contains math or drawing: {prefix}")
        old_docx = visible_text(paragraph)
        if normalized(old_docx) != normalized(old_md.replace("`", "")):
            raise ValueError(f"DOCX/Markdown mismatch: {prefix}")
        replace_text(paragraph, old_docx, new.replace("`", ""))
        markdown = markdown.replace(old_md, new, 1)
        changes.append({"old": old_md, "new": new, "reason_cn": reason})
    for old, new, reason in FRAGMENTS:
        if markdown.count(old) != 1:
            raise ValueError(f"Markdown match not unique: {old}")
        doc_old, doc_new = old.removeprefix("### "), new.removeprefix("### ")
        matches = [p for p in doc.paragraphs if normalized(doc_old) in normalized(visible_text(p))]
        if len(matches) != 1:
            raise ValueError(f"DOCX match not unique: {old}")
        replace_text(matches[0], doc_old, doc_new)
        markdown = markdown.replace(old, new, 1)
        changes.append({"old": old, "new": new, "reason_cn": reason})

    old_status = next(b for b in markdown.split("\n\n") if b.startswith("> **Revision status"))
    status = "Author working draft, refined 27 September 2026 from the supplied 26 September revision. Figure 6 and Table 4 retain pre-correction contact results; corrected coupled integration remains pending."
    markdown = markdown.replace(old_status, "> " + status)
    status_paragraph = doc.paragraphs[1]
    replace_text(status_paragraph, visible_text(status_paragraph), status)
    changes.append({"old": old_status, "new": "> " + status,
                    "reason_cn": "保留工作稿及待重算标签，区分用户审稿日期与本次小修日期。"})

    docx_path = OUTPUT / "HAMS_manuscript_refined_20260927.docx"
    md_path = OUTPUT / "manuscript_refined_20260927.md"
    doc.save(docx_path)
    md_path.write_text(markdown, encoding="utf-8")
    for figure in (PACKAGE / "manuscript" / "figures").iterdir():
        if figure.is_file():
            (OUTPUT / "figures").mkdir(exist_ok=True)
            shutil.copy2(figure, OUTPUT / "figures" / figure.name)
    shutil.copy2(PACKAGE / "manuscript" / "references.bib", OUTPUT / "references.bib")
    before, after = structural_record(SOURCE_DOCX), structural_record(docx_path)
    if before != after:
        raise AssertionError("A table, equation, media object or drawing count changed")
    if sha(SOURCE_DOCX) != source_hash:
        raise AssertionError("Input manuscript modified")
    diff = "".join(difflib.unified_diff(original_md.splitlines(True), markdown.splitlines(True),
                                      fromfile="supplied_manuscript_20260926.md", tofile=md_path.name))
    (OUTPUT / "manuscript_changes.diff").write_text(diff, encoding="utf-8")
    manifest = json.loads((PACKAGE / "code_patch" / "expected_original_sha256.json").read_text(encoding="utf-8"))
    patch_check = [{"path": name, "expected_sha256": expected,
                    "actual_sha256": sha(ROOT / name) if (ROOT / name).exists() else None,
                    "matches": sha(ROOT / name) == expected if expected is not None else not (ROOT / name).exists()}
                   for name, expected in manifest.items()]
    checks = json.loads((INPUT / "independent_checks.json").read_text(encoding="utf-8"))
    report = {
        "status": "EDITORIAL_REVISION_WITH_CORRECTED_CONTACT_RERUN_PENDING",
        "source_docx": str(SOURCE_DOCX.relative_to(ROOT)), "source_docx_sha256": source_hash,
        "source_markdown_sha256": sha(SOURCE_MD), "output_docx_sha256": sha(docx_path),
        "output_markdown_sha256": sha(md_path),
        "reference_evidence": [{"path": str((INPUT / name).relative_to(ROOT)), "sha256": sha(INPUT / name)}
                               for name in ("independent_checks.json", "fresh_hams_comparison.json", "HAMS_论文审查与真实性证据.md")],
        "checks": {"source_preserved": True, "all_tables_identical": before["tables"] == after["tables"],
                   "all_omml_identical": before["math"] == after["math"],
                   "all_embedded_images_identical": before["media"] == after["media"],
                   "tables": len(after["tables"]), "omml_objects": len(after["math"]),
                   "figures": after["inline_shapes"], "candidate_original_hashes_match": all(r["matches"] for r in patch_check)},
        "execution_scope": "Editorial refinement, supplied-evidence cross-checks and patch applicability dry-run only. No candidate patch installed and no corrected coupled simulation executed in this revision.",
        "supplied_review_execution_environment": {key: checks[key] for key in ("python", "numpy", "platform", "review_date")},
        "candidate_patch_preflight": patch_check,
        "changes": changes,
    }
    (OUTPUT / "revision_evidence.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "checks": report["checks"], "edits": len(changes)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    refine()
