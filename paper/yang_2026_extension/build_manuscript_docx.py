from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PAPER_DIR = Path(__file__).resolve().parent
SOURCE = PAPER_DIR / "manuscript.md"
OUTPUT = PAPER_DIR / "Coupled_Hydrodynamic_FourLeg_Sea_Landing_Manuscript.docx"

INK = RGBColor(0, 0, 0)
BLUE = RGBColor(35, 92, 122)
MUTED = RGBColor(92, 101, 110)
LIGHT = "EAF0F4"


def set_font(run, name: str, size: float, *, bold: bool | None = None, italic: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        for cell, value in zip(row.cells, widths_dxa):
            set_cell_width(cell, value)


def configure_table_rows(table) -> None:
    """Repeat the header and keep individual rows intact across page breaks."""
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_index == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            tr_pr.append(repeat_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_font(run, "Arial", 8.5, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = INK
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(46)
    title.paragraph_format.space_after = Pt(24)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, color, before, after in [
        ("Heading 1", 14, INK, 13, 5),
        ("Heading 2", 12, INK, 10, 4),
        ("Heading 3", 11, INK, 8, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "COUPLED HYDRODYNAMIC AND FOUR-LEG SEA-LANDING DYNAMICS"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(header.runs[0], "Arial", 7.8, color=INK)
    add_page_number(section.footer.paragraphs[0])


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("**", "")
    text = text.replace("---", "-")
    text = text.replace("\\(", "").replace("\\)", "")
    for _ in range(3):
        text = re.sub(r"\\(?:mathrm|mathbf|mathit|boldsymbol)\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\(?:mathrm|mathbf|mathit|boldsymbol)\s*", "", text)

    # Preserve common inline mathematics when Markdown is converted to Word.
    # Display equations are handled separately by add_equation().
    replacements = {
        "\\infty": "∞",
        "\\omega": "ω",
        "\\gamma": "γ",
        "\\eta": "η",
        "\\delta": "δ",
        "\\Delta": "Δ",
        "\\tau": "τ",
        "\\pi": "π",
        "\\pm": "±",
        "\\times": "×",
        "\\deg": "deg",
        "A_∞": "A∞",
        "K_h": "Kₕ",
        "F_w": "Fwave",
        "F_p": "Fplume",
        "F_c": "Fcontact",
        "H_s": "Hₛ",
        "T_p": "Tₚ",
        "m_0": "m₀",
        "S_η": "Sη",
        "_i": "ᵢ",
        "_k": "ₖ",
        "_p": "ₚ",
        "_s": "ₛ",
        "_h": "ₕ",
        "_x": "ₓ",
        "η_1": "η₁",
        "η_2": "η₂",
        "η_3": "η₃",
        "η_4": "η₄",
        "η_5": "η₅",
        "η_6": "η₆",
        "ω_k": "ωₖ",
        "^T": "ᵀ",
        "^2": "²",
        "^3": "³",
        "^6": "⁶",
        "kg/m3": "kg/m³",
        "m/s2": "m/s²",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = text.replace("\\", "")
    text = text.replace("{", "").replace("}", "")
    return text.strip()


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = False
    normalized = text.strip().lower()
    equation_leads = (
        "is written as",
        "satisfy",
        "in component form,",
        "so that",
        "vertical force is",
        "reference point,",
        "thus,",
    )
    paragraph.paragraph_format.keep_with_next = normalized.endswith(equation_leads)
    run = paragraph.add_run(clean_inline(text))
    set_font(run, "Times New Roman", 10.5)


def add_equation(doc: Document, lines: list[str], equation_number: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3.37), WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.65), WD_TAB_ALIGNMENT.RIGHT)
    raw = " ".join(line.strip() for line in lines)
    if "(\\mathbf M+\\mathbf A_\\infty)" in raw:
        equation = "(M + A∞) η̈(t) + ∫₀ᵗ Kr(t - τ) η̇(τ) dτ + Cextη̇(t) + Kₕη(t) = Fwave(t) + Fplume(t) + Fcontact(t)"
    elif "\\nabla^2\\varphi=0" in raw:
        equation = "∇²φ = 0"
    elif "c(P)\\varphi" in raw:
        equation = "c(P)φ(P) + ∫SB φ(Q) ∂G(P,Q)/∂nQ dSQ = ∫SB G(P,Q) ∂φ(Q)/∂nQ dSQ"
    elif "\\mathbf M_q(\\mathbf q)" in raw:
        equation = "Mq(q)q̈ + h(q,q̇) = Qg + Qjoint + Qabsorber + Qcontact"
    elif "\\mathbf K_r(t)=" in raw or "\\mathbf K(t)=" in raw:
        equation = "Kr(t) = (2/π) ∫₀∞ B(ω) cos(ωt) dω"
    elif "F_{b,i}=" in raw:
        equation = "Fb,i = Fs,dig(si) + Fd,dig[max(ṡi, 0)] + Fhs(si, ṡi)"
    elif "F_{n,i}=" in raw:
        equation = "Fn,i = max[0, knδi - meffGn vn,i]"
    elif "\\mathbf F_c=" in raw:
        equation = "Fcontact = -Σᵢ fᵢ;    Mcontact = -Σᵢ[(rᵢ - rR) × fᵢ]"
    elif "\\dot{\\mathbf z}_{c,k}" in raw:
        equation = "żc,k = η̇ - ωₖ zs,k;    żs,k = ωₖ zc,k"
    elif "u_x=" in raw:
        equation = "u_x = η₁ + η₅z - η₆y;    u_y = η₂ + η₆x - η₄z;    u_z = η₃ + η₄y - η₅x"
    elif "m_0=" in raw:
        equation = "m₀ = ∫ Sη(ω) dω = Hₛ²/16"
    elif "x_j(t)=" in raw:
        equation = "xⱼ(t) = Σₙ Re{Hⱼ(ωₙ,β)[2Sη(ωₙ)Δω]¹ᐟ² exp[i(ωₙt + εₙ)]}"
    elif "\\boldsymbol\\eta(t)=" in raw:
        equation = "η(t) = η_w+p(t) + Δη_c(t)"
    else:
        equation = clean_inline(raw)
    paragraph.add_run("\t")
    run = paragraph.add_run(equation)
    set_font(run, "Cambria Math", 10.5, italic=True)
    paragraph.add_run("\t")
    number_run = paragraph.add_run(f"({equation_number})")
    set_font(number_run, "Times New Roman", 10.0)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    section = doc.sections[-1]
    available = int((section.page_width - section.left_margin - section.right_margin) / 635) - 120
    first_header = clean_inline(rows[0][0]).strip()
    if columns == 5:
        if first_header == "文件":
            ratios = [0.30, 0.18, 0.24, 0.08, 0.20]
        elif first_header == "声明":
            ratios = [0.18, 0.18, 0.18, 0.38, 0.08]
        elif first_header == "图件":
            ratios = [0.26, 0.30, 0.14, 0.08, 0.22]
        else:
            ratios = [0.20, 0.24, 0.16, 0.18, 0.22]
        widths = [int(available * ratio) for ratio in ratios]
        widths[-1] += available - sum(widths)
    elif columns == 4:
        ratios = [0.09, 0.25, 0.31, 0.35] if first_header == "严重度" else [0.22, 0.29, 0.19, 0.30]
        widths = [int(available * ratio) for ratio in ratios]
        widths[-1] += available - sum(widths)
    else:
        widths = [available // columns] * columns
        widths[-1] += available - sum(widths)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = clean_inline(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_font(run, "Times New Roman", 8.2, bold=row_index == 0)
            if row_index == 0:
                shade_cell(cell, LIGHT)
    set_table_geometry(table, widths)
    configure_table_rows(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, alt: str, relative_path: str) -> None:
    image_path = PAPER_DIR / relative_path
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    width = 5.15 if relative_path.endswith("fig02-barge-geometry-mesh.png") else 6.25
    shape = run.add_picture(str(image_path), width=Inches(width))
    description = clean_inline(alt)
    shape._inline.docPr.set("descr", description)
    shape._inline.docPr.set("title", description)


def new_numbering_id(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    style_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    source = numbering.find(f".//{{{qn('w:num').split('}')[0][1:]}}}num[@{qn('w:numId')}='{style_num_id}']")
    if source is None:
        abstract_id = 0
    else:
        abstract_id = int(source.find(qn("w:abstractNumId")).get(qn("w:val")))
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def parse_bib_entries() -> list[str]:
    text = (PAPER_DIR / "references.bib").read_text(encoding="utf-8")
    entries = []
    for block in re.findall(r"@\w+\{[^,]+,(.*?)\n\}", text, flags=re.S):
        fields = {
            key.lower(): value.strip().strip("{}").replace("---", " - ")
            for key, value in re.findall(r"(\w+)\s*=\s*\{(.*?)\}\s*,?", block, flags=re.S)
        }
        authors = fields.get("author", "").replace(" and ", "; ")
        venue = fields.get("journal") or fields.get("institution", "")
        item = f"{authors}. {fields.get('title', '')}. {venue}, {fields.get('year', '')}"
        if fields.get("volume"):
            item += f", {fields['volume']}"
        if fields.get("number"):
            item += f"({fields['number']})"
        if fields.get("pages"):
            item += f": {fields['pages']}"
        item += "."
        if fields.get("doi"):
            item += f" https://doi.org/{fields['doi']}"
        entries.append(item)
    return entries


def build() -> Path:
    doc = Document()
    configure_document(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()

    title = clean_inline(lines[0].lstrip("# "))
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(46)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(title)
    set_font(run, "Times New Roman", 18, bold=True, color=INK)
    p.paragraph_format.space_after = Pt(24)

    index = 1
    paragraph_buffer: list[str] = []
    numbered_list_id: int | None = None
    equation_number = 0

    def flush_paragraph() -> None:
        if paragraph_buffer:
            add_body_paragraph(doc, " ".join(paragraph_buffer))
            paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            numbered_list_id = None
            index += 1
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            heading = clean_inline(stripped[3:])
            if heading == "References":
                doc.add_page_break()
            doc.add_paragraph(heading, style="Heading 1")
            if heading == "References":
                reference_num_id = new_numbering_id(doc)
                for number, entry in enumerate(parse_bib_entries(), start=1):
                    p = doc.add_paragraph(style="List Number")
                    apply_numbering(p, reference_num_id)
                    p.paragraph_format.space_after = Pt(3)
                    run = p.add_run(entry)
                    set_font(run, "Times New Roman", 9)
                break
            index += 1
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(clean_inline(stripped[4:]), style="Heading 2")
            index += 1
            continue
        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            flush_paragraph()
            add_figure(doc, image_match.group(1), image_match.group(2))
            index += 1
            continue
        if stripped == "\\[":
            flush_paragraph()
            equation_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != "\\]":
                equation_lines.append(lines[index])
                index += 1
            equation_number += 1
            add_equation(doc, equation_lines, equation_number)
            index += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            if len(rows) > 1 and all(set(cell) <= {"-", ":"} for cell in rows[1]):
                rows.pop(1)
            add_markdown_table(doc, rows)
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if numbered:
            flush_paragraph()
            if numbered_list_id is None:
                numbered_list_id = new_numbering_id(doc)
            p = doc.add_paragraph(style="List Number")
            apply_numbering(p, numbered_list_id)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(clean_inline(numbered.group(2)))
            set_font(run, "Times New Roman", 10.3)
            index += 1
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            numbered_list_id = None
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(clean_inline(stripped[2:]))
            set_font(run, "Times New Roman", 10.3)
            index += 1
            continue
        paragraph_buffer.append(stripped)
        numbered_list_id = None
        index += 1

    flush_paragraph()
    properties = doc.core_properties
    properties.title = title
    properties.subject = "Partitioned hydro-mechanical framework for four-leg sea landing"
    properties.author = "Sea-landing dynamics research project"
    properties.keywords = "potential-flow hydrodynamics, radiation memory, multibody contact, reusable launch vehicle, sea landing"
    doc.save(OUTPUT)
    print(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    build()
