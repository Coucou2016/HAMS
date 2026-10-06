"""Publish corrected contact results from completed, independently retained runs."""
from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import shutil
import sys
from pathlib import Path

import numpy as np
from docx import Document
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
from refine_review_20260927 import visible_text, replace_text, normalized, structural_record
from analysis.rocket_recovery.yang_2026_figures import configure_style
import matplotlib.pyplot as plt


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_figure(report, response, raw, output):
    configure_style()
    plt.rcParams.update({'font.size': 10.5, 'axes.labelsize': 10.5, 'xtick.labelsize': 9.5,
                         'ytick.labelsize': 9.5, 'legend.fontsize': 9, 'axes.titlesize': 10.5,
                         'pdf.fonttype': 42, 'ps.fonttype': 42})
    fig, ax = plt.subplots(4, 2, figsize=(7.15, 9.1), layout='constrained')
    t = np.asarray(raw['time_s']) - 510
    mask = t >= -.5
    colors = ['#0072B2', '#D55E00', '#009E73', '#CC79A7']
    for i, (leg, color) in enumerate(zip(sorted(raw['forces']['leg_contact_force_n']), colors)):
        ax[0, 0].plot(t[mask], np.asarray(raw['forces']['leg_contact_force_n'][leg])[mask]/1e6, color=color, label=leg.replace('_', ' '))
        ax[0, 1].plot(t[mask], np.asarray(raw['contact']['leg_stroke_m'][leg])[mask]*1e3, color=color)
        ax[3, 0].step(t[mask], np.asarray(raw['contact']['leg_contact'][leg])[mask] + 1.3*i, where='post', color=color)
    ax[0, 0].set(title='(a) Positive vertical foot force', ylabel='Force (MN)')
    ax[0, 0].legend(ncol=2)
    ax[0, 1].set(title='(b) Main-buffer compression', ylabel='Stroke (mm)')
    for axis in ax[0, :]:
        axis.set_xlim(-.1, 4.)
    tp = np.asarray(response['platform_time_s']) - 510
    base, coupled = np.asarray(response['baseline_q_active']), np.asarray(response['final_q_active'])
    for j, channel, scale, label in [(0, 0, 1., 'Heave (m)'), (1, 1, 180/np.pi, 'Roll (deg)')]:
        ax[1, j].plot(tp, base[:, channel]*scale, '--', color='#555555', label='No-contact baseline')
        ax[1, j].plot(tp, coupled[:, channel]*scale, color='#0072B2', label='Updated platform')
        ax[1, j].set(title=f'({chr(99+j)}) Platform response', ylabel=label)
        if j == 0:
            ax[1, j].legend(loc='lower right')
    passes = report['iterative_run']['passes']
    ax[2, 0].plot([p['pass'] for p in passes], [p['contact_summary']['max_leg_contact_force_kn']/1000 for p in passes], 'o-', color='#0072B2')
    ax[2, 0].set(title='(e) Interface iterations', xlabel='Pass', ylabel='Peak force (MN)', xticks=[p['pass'] for p in passes])
    v = report['time_step_convergence']['values']
    names = ['dt', 'dt_over_2', 'dt_over_4']
    ax[2, 1].plot([.5, .25, .125], [v[k]['max_leg_contact_force_kn']/1000 for k in names], 'o-', color='#0072B2', label='Peak force')
    twin = ax[2, 1].twinx()
    twin.plot([.5, .25, .125], [v[k]['max_leg_stroke_m']*1000 for k in names], 's--', color='#D55E00', label='Stroke')
    ax[2, 1].set(title='(f) Contact-step sensitivity', xlabel='Contact step (ms)', ylabel='Peak force (MN)', xticks=[.5,.25,.125])
    ax[2, 1].invert_xaxis()
    twin.set_ylabel('Peak stroke (mm)', color='#D55E00')
    ax[3, 0].set(title='(g) Four-foot contact states', ylabel='Foot', yticks=[.5+1.3*i for i in range(4)], yticklabels=['1','2','3','4'])
    rows = report['contact_regularization_sensitivity']['rows']
    ax[3, 1].plot([r['normal_stiffness_n_m']/1e6 for r in rows], [r['peak_contact_force_kn']/1000 for r in rows], 'o-', color='#0072B2')
    ax[3, 1].set(title='(h) Contact-stiffness sensitivity', xlabel='Normal stiffness (MN/m)', ylabel='Peak force (MN)')
    for axis in [*ax[0, :], *ax[1, :], ax[3, 0]]:
        axis.set_xlabel('Time from nominal touchdown (s)')
    output.parent.mkdir(parents=True, exist_ok=True)
    for suffix in ('.png', '.svg', '.pdf'):
        fig.savefig(output.with_suffix(suffix), dpi=600, bbox_inches='tight')
    plt.close(fig)


def publish(case, output):
    source = HERE / 'revision_20260927'
    report_file = case / 'chrono-same-platform-multibody-report.json'
    report = read(report_file)
    response_file = case / 'Output/RocketRecovery/chrono-same-platform-multibody-response.json'
    response = read(response_file)
    audit = read(case / 'corrected-run-audit.json')
    expected_histories = len(list((case / 'raw').glob('*/*.audit.json')))
    if not (audit['completed_report_exists'] and audit['all_histories_complete']
            and audit['impulse_transfer_pass'] and len(audit['histories']) == expected_histories):
        raise ValueError('Full integration histories must pass the audit before publication')
    run = report['iterative_run']
    raw_file = Path(run['passes'][-1]['raw_history']['path'])
    raw = read(raw_file)
    output.mkdir(parents=True, exist_ok=True)
    shutil.copytree(source / 'figures', output / 'figures', dirs_exist_ok=True)
    shutil.copy2(source / 'references.bib', output / 'references.bib')
    figure = output / 'figures/fig06-partitioned-contact-feedback.png'
    make_figure(report, response, raw, figure)
    doc = Document(source / 'HAMS_manuscript_refined_20260927.docx')
    original = (source / 'manuscript_refined_20260927.md').read_text(encoding='utf-8')
    md = original
    changes = []

    def fragment(old, new):
        nonlocal md
        if md.count(old) != 1:
            raise ValueError(f'Markdown fragment not unique: {old}')
        matches = [p for p in doc.paragraphs if normalized(old) in normalized(visible_text(p))]
        if len(matches) != 1:
            raise ValueError(f'Word fragment not unique: {old}')
        replace_text(matches[0], old, new)
        md = md.replace(old, new, 1)
        changes.append({'old': old, 'new': new})

    def paragraph(prefix, new):
        nonlocal md
        matches = [p for p in doc.paragraphs if normalized(visible_text(p)).startswith(prefix)]
        blocks = [b for b in md.split('\n\n') if b.startswith(prefix)]
        if len(matches) != 1 or len(blocks) != 1:
            raise ValueError(f'Paragraph not unique: {prefix}')
        p = matches[0]
        if p._p.xpath('.//m:oMath | .//w:drawing'):
            raise ValueError(f'Cannot replace native math paragraph: {prefix}')
        old = visible_text(p)
        replace_text(p, old, new)
        md = md.replace(blocks[0], new, 1)
        changes.append({'old': blocks[0], 'new': new})

    count = len(run['passes'])
    s = run['final_contact_summary']
    c = report['coupling_iteration_convergence']
    last = c['comparisons'][-1]
    interface = 'meets' if c['pass'] else 'does not meet'
    all_closed = 'All six distinct runs meet' if report['all_interface_runs_closed'] else 'Not all six distinct runs meet'
    step_status = 'meets' if report['time_step_convergence']['pass'] else 'does not meet'
    force = s['max_leg_contact_force_kn']/1000
    stroke = s['max_leg_stroke_m']*1000
    transfer_error = max(r['transfer_max_abs_residual'] for r in audit['histories'])
    headline = (f'The corrected full-scale calculation {interface} the 2% interface update criterion after {count} passes. '
                f'Its peak positive vertical foot force is {force:.3f} MN and peak buffer stroke is {stroke:.2f} mm. '
                f'The contact-step study {step_status} the joint 5% criterion. '
                'These are newly integrated responses of the stated proxy model, not validated full-scale landing loads.')
    fragment('Author working draft, refined 27 September 2026 from the supplied 26 September revision. Figure 6 and Table 4 retain pre-correction contact results; corrected coupled integration remains pending.',
             'Corrected computational revision, 28 September 2026. Figure 6 and Table 4 use new coupled-contact integrations. Numerical qualification and physical validation are distinguished explicitly.')
    fragment('The archived full-scale calculation meets a 2% fixed-point update criterion after four passes but fails contact-refinement checks. Its inertia mapping, collision-deck reference and executed load transfer also require correction. The corresponding landing histories are retained solely as pre-correction diagnostics. Independent subproblem reruns support computational reproducibility within the stated scope; corrected coupled integration and renewed convergence tests are required before interpreting the full-scale contact loads or touchdown sequence physically.', headline)
    fragment('The full-scale example is retained as a diagnostic pending the corrections described in Sections 2.4-2.5.',
             'The full-scale example uses the coordinate-consistent and impulse-preserving implementation described in Sections 2.4-2.5; its numerical limitations are assessed separately from the reduced drop-test comparison.')
    fragment('All archived full-scale contact histories in Section 4.4 precede this correction. The candidate implementation also closes mass, center of gravity and inertia',
             'The histories reported in Section 4.4 have been recomputed with the corrected mapping. The implementation closes mass, center of gravity and inertia')
    fragment('The candidate recorder stores', 'The integration-grid recorder stores')
    paragraph('Each foot has its own contact state.',
              'Each foot has its own contact state. A sample is classified as contact when positive vertical force exceeds 100 N or normal penetration exceeds 10 micrometers. First contact, lift-off and re-contact are detected from transitions of this indicator, so event times inherit the contact step and these thresholds. Gravity activation separately uses the first force-threshold crossing. All recorded integration nodes are retained for peak and event audits; compact plotting histories are separate derived files. The event indicator and deck-normal force projection are distinct quantities.')
    paragraph('The archived main calculation uses pointwise',
              'The fine-grid wrench is integrated as a piecewise-linear function over intervals bounded by target-time midpoints. Each interval integral is divided by its width. With identical source and target endpoints, the half-width end intervals make the target-grid trapezoidal impulse equal to the fine-grid impulse in all six components. This operation replaces pointwise interpolation in the executed feedback path. It preserves resultant impulse, not force peaks or interface work; contact-step and platform-response accuracy require separate checks.')
    fragment('The archived main case uses four prescribed passes; the last update is then checked against a 2% criterion for the platform trajectory, peak vertical force, stroke, touchdown span and positive vertical impulse. This post-run check should not be described as an implemented adaptive stopping rule. The candidate runner applies the same criterion after at least two passes, up to eight, and reports failure when the limit is reached without closure.',
             'The same 2% update criterion is applied to the platform trajectory, peak positive vertical force, stroke, touchdown span and positive vertical impulse. Iteration stops after at least two passes when all criteria are met, with an eight-pass limit. Reaching that limit without closure is reported as failure. Production, time-step and stiffness cases use this identical stopping rule.')
    fragment('Figure 1 depicts the intended interface rather than certifying the archived implementation.',
             'Figure 1 depicts the implemented motion and resultant-load interface, not an independent contact-energy validation.')
    paragraph('The contact archive does not use a geometrically identical collision deck:',
              'The collision deck reads its 120 m by 50 m dimensions and 3 m elevation from the hydrodynamic-platform configuration. One rotation convention defines its collision pose, surface normal, local point coordinates and spatial angular velocity. The deck-center velocity includes the rotational offset from the platform reference point. Contact moments are evaluated about the instantaneous platform reference rather than an unrelated fixed world origin. The pre-correction archive is retained separately for traceability.')
    fragment('The archived landing example uses', 'The corrected landing example uses')
    fragment('The archived production interval is', 'The production interval is')
    paragraph('The archived production calculation has four passes.',
              f'The production calculation uses {count} interface passes over 506-526 s. The contact-step study uses 0.0005, 0.00025 and 0.000125 s over 506-516 s. Contact-stiffness cases use 100, 250 and 500 MN/m at 0.000125 s over that same shorter interval. All runs use the identical adaptive 2% interface criterion, and every recorded contact-grid node is retained. The 100 MN/m stiffness case reuses the identical finest-step calculation rather than rerunning it. {all_closed} the interface criterion; physical parameters are not adjusted to force acceptance.')
    fragment('That separate review environment lacked PyChrono, and no corrected full-scale contact history is supplied.',
             'The separate review environment lacked PyChrono; the corrected integrations reported here were subsequently executed locally with the recorded Project Chrono 10.0.0 environment.')
    fragment('4.4 Archived sequential touchdown and platform feedback', '4.4 Corrected sequential touchdown and platform feedback')
    paragraph('Figure 6 and Table 4 document',
              'Figure 6 and Table 4 are regenerated from the corrected coupled calculations. Wave, plume and contact forcing use the same platform operator; the contact model uses the corrected inertia allocation, collision-deck reference and impulse-preserving load transfer. The original uncorrected histories are preserved outside this result set. The new calculations resolve the implementation defects without removing the separate limitations of the proxy mechanism, finite-startup wave state or constrained horizontal platform coordinates.')
    paragraph('Within that implementation,',
              f'The last production update has a platform fixed-point metric of {100*last["platform_fixed_point_metric"]:.4f}%. The largest relative change among peak force, stroke, touchdown span and positive vertical impulse is {100*max(last["relative_changes"].values()):.4f}%. Thus this run {interface} the selected 2% interface criterion after {count} passes. Interface closure tests repeatability of the partitioned trajectory and selected outputs; it does not establish time-step convergence or experimental accuracy.')
    # Captions are text fragments, preserving their existing Word formatting.
    cap_old = next(b for b in md.split('\n\n') if b.startswith('**Figure 6.**'))
    cap_new = ('Figure 6. Corrected four-foot calculation using the same-platform operator. Force is the positive world-vertical component, not the deck-normal projection. '
               f'The production traces use {count} adaptive passes over 506-526 s and all recorded 0.5 ms nodes. '
               'Force and stroke panels resolve the impact window; platform and contact-state panels show the continued response. Step and stiffness comparisons use 506-516 s with the same interface criterion. Normal-force projections are retained separately in the raw records.')
    caption = next(p for p in doc.paragraphs if visible_text(p).startswith('Figure 6.'))
    replace_text(caption, visible_text(caption), cap_new)
    md = md.replace(cap_old, '**Figure 6.** ' + cap_new.split('. ', 1)[1])
    state = s['contact_state']
    sequence = ', '.join(x['leg_id'].replace('leg_', '') for x in state['touchdown_sequence'])
    platform_summary = run['final_platform_summary']
    paragraph('The archived production sequence is',
              f'The corrected initial-contact order is legs {sequence}, spanning {state["touchdown_span_s"]:.4f} s. Peak platform heave, roll and pitch are {platform_summary["heave_peak_m"]:.4f} m, {platform_summary["roll_peak_deg"]:.4f} deg and {platform_summary["pitch_peak_deg"]:.4f} deg. The assembled equal-and-opposite wrench residual remains an algebraic bookkeeping check because both arrays use the same foot-force samples. Surge, sway and yaw contact loads are calculated but not applied to platform motion; their omission is a model constraint, not evidence that these loads vanish.')
    paragraph('The archived platform work ledger has',
              'The platform work ledger checks the discrete platform equation under its applied input. It does not close vehicle propulsion, foot-contact elasticity, articulated-body kinetic energy and platform motion into one conserved system. The former mixed-origin contact-pair work proxy is disabled and stored as unavailable, rather than assigned a zero residual. An independent contact-pair energy balance remains outside the available evidence.')
    vals = report['time_step_convergence']['values']
    metrics = report['time_step_convergence']['pairwise'][-1]['relative_changes']
    step_closure = []
    from analysis.rocket_recovery.chrono_same_platform_multibody import coupling_iteration_summary
    for dt in (.0005, .00025, .000125):
        step_run = read(case / 'raw' / f'end516_dt{dt:g}_kn1e+08' / 'run-result.json')
        closes = coupling_iteration_summary(step_run['passes'])['pass']
        step_closure.append(f'{dt*1000:g} ms: {len(step_run["passes"])} passes, ' + ('closed' if closes else 'not closed'))
    paragraph('The archived step study does not qualify',
              f'The corrected contact-step study {step_status} its joint 5% criterion (Table 4). Interface outcomes are ' + '; '.join(step_closure) + f'. From 0.25 to 0.125 ms, peak positive vertical force changes by {100*metrics["max_leg_contact_force_kn"]:.2f}%, buffer stroke by {100*metrics["max_leg_stroke_m"]:.2f}%, penetration by {100*metrics["max_contact_penetration_m"]:.2f}% and initial touchdown span by {100*metrics["touchdown_span_s"]:.2f}%. Positive vertical impulse changes by {100*metrics["total_normal_impulse_ns"]:.3f}%. A small integral change does not bound peak or event-time error. These comparisons use matched 506-516 s intervals and the same stopping rule. Where interface closure fails, their differences mix iteration and time-step effects and cannot establish isolated time-step convergence.')
    table = doc.tables[3]
    table_rows = []
    for i, (key, dt) in enumerate(zip(['dt','dt_over_2','dt_over_4'], [.5,.25,.125]), 1):
        r = vals[key]
        cells = [f'{dt:g} ms', f'{r["max_leg_contact_force_kn"]/1000:.3f} MN', f'{r["max_leg_stroke_m"]*1000:.2f} mm', f'{r["max_contact_penetration_m"]*1000:.2f} mm', f'{r["touchdown_span_s"]:.4f} s']
        for cell, value in zip(table.rows[i].cells, cells):
            p = cell.paragraphs[0]
            if p.runs:
                p.runs[0].text = value
                for rr in p.runs[1:]: rr.text = ''
            else: p.add_run(value)
        table_rows.append('| ' + ' | '.join(cells) + ' |')
    old_table = next(b for b in md.split('\n\n') if b.startswith('| Contact step |'))
    md = md.replace(old_table, '\n'.join(old_table.splitlines()[:2] + table_rows))
    old_cap = next(b for b in md.split('\n\n') if b.startswith('**Table 4.**'))
    new_cap = 'Table 4. Corrected contact-step study over 506-516 s with a common adaptive interface criterion. Peak force is the positive world-vertical foot force; convergence is assessed jointly across all registered metrics.'
    p = next(p for p in doc.paragraphs if visible_text(p).startswith('Table 4.'))
    replace_text(p, visible_text(p), new_cap)
    md = md.replace(old_cap, '**Table 4.** ' + new_cap.split('. ', 1)[1])
    reg = report['contact_regularization_sensitivity']['rows']
    fragment('The one-pass stiffness study gives a separate regularization sensitivity.',
             'The corrected stiffness study evaluates contact regularization using the same adaptive interface rule.')
    fragment('from 100 to 250 and 500 MN/m changes peak positive vertical force from 6.432 to 8.718 and 8.999 MN, and stroke from 73.60 to 84.91 and 94.16 mm. Penetration decreases from 52.42 to 23.64 and 12.65 mm. The corresponding positive vertical impulse is approximately 3.909 MN s. Its relative stability within this study does not resolve the different interface iteration count or provide a corrected physical impulse. In particular, one should not combine the 100 MN/m one-pass value with a different-pass time-step row as though they were the same run.',
             'from 100 to 250 and 500 MN/m gives peak positive vertical forces of ' + ', '.join(f'{r["peak_contact_force_kn"]/1000:.3f}' for r in reg) + ' MN, and buffer strokes of ' + ', '.join(f'{r["maximum_buffer_stroke_m"]*1000:.2f}' for r in reg) + ' mm. Maximum penetrations are ' + ', '.join(f'{r["maximum_penetration_m"]*1000:.2f}' for r in reg) + ' mm. These are numerical sensitivities of the stated penalty law, not measured material behavior. The 100 MN/m entry is exactly the finest-step run in Table 4. Any run that reaches the eight-pass limit without interface closure remains unqualified; increasing stiffness alone does not establish an accurate contact load.')
    standing = s['stable_standing_diagnostic']
    paragraph('At 526 s, the archived calculation',
              f'At 526 s, the corrected calculation has {s["final_contact_count"]} feet in contact. Its final world-vertical speed is {s["final_vertical_velocity_m_s"]:.4f} m/s and angular-rate magnitude is {standing["final_angular_rate_deg_s"]:.4f} deg/s. The selected standing-diagnostic flag is {str(standing["all_flags_true"]).lower()}. These checks use prescribed 0.05 m/s and 0.25 deg/s limits and are not a measured landing-success criterion. In particular, absolute vehicle velocity is not a substitute for relative vehicle-deck settling on a moving platform.')
    paragraph('The stored multibody response contains',
              f'The production history retains {len(raw["time_s"]):,} actual integration-grid records per pass, including both endpoints. All {len(audit["histories"])} retained histories pass sample-count, time-spacing and hash checks. The largest absolute source-to-target impulse residual across six components is {transfer_error:.3g} in the corresponding N s or N m s units. This is a quadrature-transfer check, not a contact-work balance. Figures use the retained raw history; compact animation records are derived products and are not substituted for peak or event audits.')
    paragraph('A corrected assessment requires',
              'The corrected calculations now use consistent inertia axes and deck references, conservative resultant-load transfer and complete recorded contact histories. Their contact-refinement and interface outcomes are reported without suppressing failed metrics. Wave-screening and reduced-test comparisons retain their separate evidence and limitations; neither validates this full-scale mechanism. Remaining numerical or physical limitations are not resolved merely by correcting implementation errors.')
    fragment('Its original inertia mapping and deck datum are inconsistent, and its reported transfer method differs from the executed code. Contact refinement is incomplete,',
             'The inertia mapping, deck datum and executed transfer path have been corrected and rerun. Contact numerical qualification remains incomplete,')
    paragraph('The archived landing meets its final update criterion', headline + ' Independent coupled experiments, measured mechanism properties and a wider numerical assessment remain necessary before design or operational use.')
    paragraph('The project repository and accompanying review package contain',
              'The project contains case inputs, meshes, locally generated hydrodynamic outputs, deck-point response operators, realization-level wave statistics, digitized comparison targets and figure-generation code. Corrected contact results are stored separately from the pre-correction archive, with per-pass integration histories, source and runtime identities, file hashes, impulse-transfer audits and refinement results. File identity and reproducibility do not establish experimental authenticity or physical accuracy.')
    fragment('The absence of PyChrono in the separate review environment limits that assessment, not the availability of the source code. No corrected coupled history or renewed contact-convergence evidence is included in this manuscript version.',
             'Corrected coupled histories and renewed contact-refinement results are included in this revision. The retained environment manifest and per-pass hashes identify the local execution and generated data; unsuccessful qualification checks remain explicitly reported.')
    # Numerical setting/evidence in Table 2 must not retain the old four-pass status.
    replacements = [f'0.0005 s production contact step; 0.01 s platform step; {count} adaptive passes',
                    'Corrected coupled calculation; numerical qualification reported separately']
    for cell, value in zip(doc.tables[1].rows[-1].cells[2:], replacements):
        cell.text = value
    md = md.replace('0.0005 s production contact step; 0.01 s platform step; four passes', replacements[0])
    md = md.replace('Archived pre-correction diagnostic; corrected rerun pending', replacements[1])
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            if 'pending' in p.text.lower(): p.text = 'Corrected computational revision | Validation limits retained'
    # Replace Figure 6 media in-place, leaving all other drawings and OMML intact.
    image = doc.inline_shapes[5]
    rid = image._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[rid]._blob = figure.read_bytes()
    with Image.open(figure) as bitmap:
        image.height = round(image.width * bitmap.height / bitmap.width)
    target_docx = output / 'HAMS_manuscript_corrected_20260928.docx'
    doc.save(target_docx)
    (output / 'manuscript_corrected_20260928.md').write_text(md, encoding='utf-8')
    before, after = structural_record(source / 'HAMS_manuscript_refined_20260927.docx'), structural_record(target_docx)
    assert before['math'] == after['math'], 'Native equations changed unexpectedly'
    assert before['tables'][0] == after['tables'][0] and before['tables'][2] == after['tables'][2]
    evidence = {'input_report_sha256': digest(report_file), 'input_response_sha256': digest(response_file),
                'raw_figure_source_sha256': digest(raw_file), 'docx_sha256': digest(target_docx),
                'native_equations_preserved': True, 'table4_source': 'time_step_convergence.values',
                'figure6_source': str(raw_file.relative_to(ROOT)), 'changes': changes,
                'all_interface_runs_closed': report['all_interface_runs_closed'],
                'contact_step_pass': report['time_step_convergence']['pass']}
    (output / 'publication-evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    (output / 'manuscript_changes.diff').write_text(''.join(difflib.unified_diff(original.splitlines(True), md.splitlines(True))), encoding='utf-8')
    note = [
        '# 修正后接触重算与证据审查', '',
        '本文件记录本机实际执行结果，不将代码修改、数值收敛和物理验证混为一谈。', '',
        '## 已落实的实现修正', '',
        '- 惯量从文献 Y 向上坐标转换到 Z 向上；总惯量为 diag(2.57e7, 2.57e7, 3.76e5) kg m²。',
        '- 按参考构型扣除四个显式脚垫的质量、质心和惯量贡献；这些部件仍是公开参数不足时声明的代理配置，不是 CAD 实测。',
        '- 碰撞甲板使用水动力配置的 120×50 m 和静水面以上 3 m；统一姿态、角速度及移动参考点力矩。',
        '- 额外修正接触阻尼换算中残留的旧甲板尺寸：实际碰撞体质量为 2,100,000 kg，换算与本机 Chrono 对象一致。',
        '- 实际六分量传力调用改为区间积分平均，检查同覆盖区间的冲量残差。',
        '- 世界坐标正向 Fz 与甲板法向投影分列保存；旧字段未冒充法向力。',
        '- 每轮保存全部已记录积分节点和哈希；细步长浮点累加漏记问题改为按记录索引计算输出时刻。',
        '- 取消不成立的混合参考点功率/能量代理，缺失项保持 unavailable/null。',
        '- 全部工况采用相同 2% 自适应反馈准则，最少两轮、最多八轮；未满足即报告失败。', '',
        '## 实际重算', '',
        '- 实际 PyChrono 环境的回归测试：76 项通过；测试不等于完整物理验证。',
        f'- 主工况：506–526 s、0.5 ms，执行 {count} 轮，接口闭合：{c["pass"]}。',
        f'- 所有独立工况接口均闭合：{report["all_interface_runs_closed"]}。',
        f'- 接触步长联合 5% 判据通过：{report["time_step_convergence"]["pass"]}。',
        f'- 主工况峰值正向竖直力 {force:.6f} MN，最大缓冲行程 {stroke:.6f} mm。',
        f'- {len(audit["histories"])} 份原始历史通过哈希、节点数量及间隔检查。', '',
        '| 步长 | 峰值正向 Fz | 峰值行程 | 最大穿透 | 初始触地跨度 |',
        '|---|---:|---:|---:|---:|', *table_rows, '',
        '## 证据边界', '',
        '图6直接使用最终轮原始节点，表4来自同一输出目录的时间步长报告。',
        '作用—反作用零残差仍只是由相同输入构造的代数检查；没有将它称为独立物理验证。',
        '传力冲量守恒不等于接触功守恒，不证明峰值准确；全部收敛失败项保留。',
        '补充的短时接触雅可比开关对照未消除步长敏感性，未据此替换生产结果或宣称优化成功。',
        '早期中断的修正输出和诊断试算保留在其他目录，不混入本版论文；本版数据源只取 _v2 目录。',
        '平台水平自由度仍受约束；波浪从有限起始时刻的零状态开始；喷流仅是平台外载。',
        '没有新增实测支腿参数或全尺度试验，未解决的水动力认证差异和 Yang 波形偏差仍保留在正文。',
        '本轮未改写 HAMS 核心，没有使用预编译 Bin 替代原有水动力输出，也没有重算并冒称新的水动力认证。', '',
        '## 复现入口', '',
        '`powershell -ExecutionPolicy Bypass -File local-tools/Run-CorrectedContact.ps1 -OutputDir RocketRecoveryCases/新独立目录`', '',
        '原始历史、环境、测试和输出审查：`' + str(case.relative_to(ROOT)) + '`。',
        '论文表图来源及文件哈希：`publication-evidence.json`。',
    ]
    (output / '修正重算审查.md').write_text('\n'.join(note) + '\n', encoding='utf-8')
    print(target_docx)


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('--case', type=Path, default=ROOT/'RocketRecoveryCases/Chrono_LeggedRecovery_Corrected20260927_v2')
    p.add_argument('--output', type=Path, default=HERE/'revision_corrected_20260928')
    a = p.parse_args()
    publish(a.case.resolve(), a.output.resolve())
